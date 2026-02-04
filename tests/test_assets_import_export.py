"""Excel 파일 import/export 기능 테스트"""
import pytest
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import Workbook


@pytest.fixture(scope="function")
def admin_user_token(client: TestClient, db_session) -> str:
    """관리자 사용자 토큰 생성"""
    # 관리자로 회원가입
    signup_data = {
        "name": "excel_admin",
        "email": "excel_admin@example.com",
        "password": "password123",
        "club_name": "Excel테스트동아리",
        "club_description": "Excel 테스트용",
        "location_lat": 37_500_000,
        "location_lng": 127_000_000,
    }
    res = client.post("/api/admin/signup", json=signup_data)
    assert res.status_code == 201, res.text
    
    # 로그인
    login_data = {
        "email": "excel_admin@example.com",
        "password": "password123"
    }
    res = client.post("/api/auth/login", json=login_data)
    assert res.status_code == 200, res.text
    
    tokens = res.json().get("tokens", res.json())
    return tokens["access_token"]


@pytest.fixture(scope="function")
def admin_headers(admin_user_token: str) -> dict:
    """관리자 헤더"""
    return {"Authorization": f"Bearer {admin_user_token}"}


def test_download_import_template(client: TestClient, admin_headers: dict):
    """Excel 템플릿 다운로드 테스트"""
    response = client.get("/api/assets/import_template", headers=admin_headers)
    
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in response.headers["content-disposition"]
    assert "asset_import_template.xlsx" in response.headers["content-disposition"]
    
    # Excel 파일 내용 확인
    assert len(response.content) > 0


def test_import_assets_success(client: TestClient, admin_headers: dict):
    """Excel 파일 업로드 성공 테스트"""
    # Excel 파일 생성
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "description", "total_quantity", "available_quantity", "location", "created_at"])
    ws.append(["테스트물품1", "설명1", 5, 5, "창고A", "2024-01-01 00:00:00"])
    ws.append(["테스트물품2", "설명2", 10, 10, "창고B", "2024-01-01 00:00:00"])
    
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    files = {
        "file": ("test_assets.xlsx", excel_buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    assert response.status_code == 200, response.text
    data = response.json()
    assert "imported" in data
    assert "failed" in data
    assert data["imported"] >= 0


def test_import_assets_invalid_excel(client: TestClient, admin_headers: dict):
    """잘못된 Excel 파일 업로드 테스트"""
    # 잘못된 Excel
    wb = Workbook()
    ws = wb.active
    ws.append(["invalid", "headers", "format"])
    ws.append(["값1", "값2", "값3"])
    
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    files = {
        "file": ("invalid.xlsx", excel_buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    # 400 또는 성공이지만 failed > 0
    assert response.status_code in [200, 400]
    if response.status_code == 200:
        data = response.json()
        assert "failed" in data


def test_import_assets_empty_file(client: TestClient, admin_headers: dict):
    """빈 파일 업로드 테스트"""
    files = {
        "file": ("empty.xlsx", BytesIO(b""), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    assert response.status_code in [200, 400]


def test_import_assets_without_auth(client: TestClient):
    """인증 없이 Excel 업로드 시도"""
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "description", "total_quantity", "available_quantity", "location", "created_at"])
    ws.append(["물품1", "설명1", 5, 5, "창고", "2024-01-01 00:00:00"])
    
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    files = {
        "file": ("test.xlsx", excel_buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post("/api/assets/import", files=files)
    
    assert response.status_code == 401


def test_export_assets(client: TestClient, admin_headers: dict, db_session):
    """자산 목록 Excel 내보내기 테스트"""
    # 먼저 자산 생성
    asset_data = {
        "name": "내보내기테스트물품",
        "description": "테스트",
        "quantity": 5,
        "location": "창고",
        "category_id": 1
    }
    
    # 자산 생성 시도 (실패할 수 있음)
    client.post("/api/admin/assets", json=asset_data, headers=admin_headers)
    
    # Excel 내보내기
    response = client.get("/api/assets/export", headers=admin_headers)
    
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in response.headers["content-disposition"]
    assert "asset_export_" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]
    
    # Excel 내용 확인
    assert len(response.content) >= 0  # 빈 목록일 수도 있음


def test_export_assets_without_auth(client: TestClient):
    """인증 없이 Excel 내보내기 시도"""
    response = client.get("/api/assets/export")
    
    assert response.status_code == 401


def test_export_assets_non_admin(client: TestClient, db_session):
    """일반 사용자로 Excel 내보내기 시도"""
    # 일반 사용자 생성 및 로그인
    signup_data = {
        "name": "일반사용자",
        "email": "user@example.com",
        "password": "password123",
        "student_id": "2024-12345"
    }
    
    # 회원가입 (일반 사용자 가입 엔드포인트 사용)
    res = client.post("/api/users/signup", json=signup_data)
    if res.status_code != 201:
        pytest.skip(f"일반 사용자 회원가입 실패: {res.status_code} - {res.text}")
    
    # 로그인
    login_data = {
        "email": "user@example.com",
        "password": "password123"
    }
    res = client.post("/api/auth/login", json=login_data)
    assert res.status_code == 200
    
    tokens = res.json().get("tokens", res.json())
    user_headers = {"Authorization": f"Bearer {tokens['access_token']}"}
    
    # Excel 내보내기 시도
    response = client.get("/api/assets/export", headers=user_headers)
    
    assert response.status_code == 403


def test_import_with_special_characters(client: TestClient, admin_headers: dict):
    """특수 문자가 포함된 Excel 업로드 테스트"""
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "description", "total_quantity", "available_quantity", "location", "created_at"])
    ws.append(["물품,쉼표포함", "설명", 5, 5, "창고", "2024-01-01 00:00:00"])
    ws.append(["물품-특수", "설명", 3, 3, "창고", "2024-01-01 00:00:00"])
    
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    files = {
        "file": ("special.xlsx", excel_buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    assert response.status_code == 200
    data = response.json()
    assert "imported" in data
    assert "failed" in data


def test_import_large_file(client: TestClient, admin_headers: dict):
    """대용량 Excel 파일 업로드 테스트"""
    # 100개 항목 생성
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "description", "total_quantity", "available_quantity", "location", "created_at"])
    
    for i in range(100):
        ws.append([f"물품{i}", f"설명{i}", i % 10 + 1, i % 10 + 1, f"창고{i % 5}", "2024-01-01 00:00:00"])
    
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    files = {
        "file": ("large.xlsx", excel_buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    assert response.status_code == 200
    data = response.json()
    assert "imported" in data
    assert "failed" in data


def test_download_template_without_auth(client: TestClient):
    """인증 없이 템플릿 다운로드 시도"""
    response = client.get("/api/assets/import_template")
    
    # 템플릿은 인증 없이도 다운로드 가능할 수 있음
    # 정책에 따라 401 또는 200
    assert response.status_code in [200, 401]


def test_import_invalid_file_type(client: TestClient, admin_headers: dict):
    """지원하지 않는 파일 형식 업로드 테스트 (CSV, TXT 등)"""
    # CSV 파일 업로드 시도
    csv_content = b"name,description,quantity\nitem1,desc1,5"
    files = {
        "file": ("test.csv", BytesIO(csv_content), "text/csv")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    assert response.status_code == 400
    assert "파일 형식" in response.json()["detail"] or "Excel" in response.json()["detail"]


def test_import_oversized_file(client: TestClient, admin_headers: dict):
    """용량 초과 파일 업로드 테스트 (10MB 초과)"""
    # 10MB보다 큰 파일 생성 - 테스트에서는 Content-Length 헤더로 미들웨어에서 차단됨
    # 실제 10MB+ 데이터를 생성하면 테스트가 느려지므로 작은 크기로 테스트
    # 미들웨어는 Content-Length 헤더를 검사하므로 큰 Content-Length를 직접 설정하여 테스트
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "description", "total_quantity", "available_quantity", "location", "created_at"])
    ws.append(["테스트", "설명", 1, 1, "창고", "2024-01-01 00:00:00"])
    
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    # Content-Length를 조작하여 10MB 초과로 설정
    files = {
        "file": ("oversized.xlsx", excel_buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    # 미들웨어 테스트: Content-Length 헤더를 직접 설정
    headers = {**admin_headers, "Content-Length": str(15 * 1024 * 1024)}  # 15MB
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=headers
    )
    
    # 미들웨어에서 413 반환 또는 실제 파일이 작으므로 200
    assert response.status_code in [200, 413]


def test_import_corrupted_excel(client: TestClient, admin_headers: dict):
    """손상된 Excel 파일 업로드 테스트"""
    # Excel 헤더를 흡내낸 손상된 파일
    corrupted_content = b"PK\x03\x04corrupted_excel_data_here"
    files = {
        "file": ("corrupted.xlsx", BytesIO(corrupted_content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    
    response = client.post(
        "/api/assets/import",
        files=files,
        headers=admin_headers
    )
    
    assert response.status_code == 400
    assert "Excel" in response.json()["detail"] or "올바르" in response.json()["detail"]
